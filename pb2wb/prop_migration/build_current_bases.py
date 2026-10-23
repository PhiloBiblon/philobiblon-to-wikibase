"""
build_current_bases.py — snapshot of the P721 bases still present in FactGrid,
annotated against the existing mapping (P721-P129.tsv) and the June snapshot.

Converted statements have their P721 qualifier removed, so every value in a
fresh SPARQL export is by definition not yet converted.

Input is the query-service export (value,freq), CSV or TSV.

Usage (from pb2wb/):
    python prop_migration/build_current_bases.py prop_migration/bases.2026-10-09.csv
    python prop_migration/build_current_bases.py EXPORT --out prop_migration/current_bases.xlsx

Status values:
    pending_run      vetted=Y with valid QID, no locator — batch should have
                     converted it (stragglers, new statements, or June batch never run)
    pending_locator  vetted=Y with valid QID, has locator — waits for --locators-only
    vetted_bad_qid   vetted=Y but QID blank/malformed
    suggested        known string, not vetted, has a QID suggestion
    unmatched        known string, no QID
    new_key_match    string not in mapping, but its parsed key matches a vetted key
    new              string not in mapping, no match
    compound         contains " / " — manual handling
    excluded         non-reference (fol. mod., fichero, …) — P73 note candidate
"""

import os
import re
import sys
import csv
import json
import argparse
from collections import Counter, defaultdict

dir_path = os.path.dirname(os.path.realpath(__file__))
parent_dir_path = os.path.dirname(dir_path)
sys.path.append(parent_dir_path)
from common.sheet_fields import is_vetted

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from prop_migration.generate_basis_mapping import preprocess, strip_html, extract_qid

MAPPING_TSV  = 'prop_migration/P721-P129.tsv'
JUNE_TSV     = 'prop_migration/P721-raw-values.tsv'
PROGRESS     = 'prop_migration/basis_migration_progress.jsonl'
OUT_XLSX     = 'prop_migration/current_bases.xlsx'

FG_ITEM = 'https://database.factgrid.de/wiki/Item:'
QID_RE = re.compile(r'^Q\d+$')

STATUS_ORDER = ['pending_run', 'pending_locator', 'vetted_bad_qid', 'suggested',
                'new_key_match', 'unmatched', 'new', 'compound', 'excluded']
STATUS_NOTE = {
    'pending_run':     'vetted, no locator — should already be converted; re-run batch',
    'pending_locator': 'vetted, has locator — waits for --locators-only pass',
    'vetted_bad_qid':  'vetted but QID blank or malformed',
    'suggested':       'known string, unvetted QID suggestion',
    'new_key_match':   'new string; parsed key matches a vetted key',
    'unmatched':       'known string, no QID found',
    'new':             'new string since June, no match',
    'compound':        'contains " / " — manual handling',
    'excluded':        'non-reference (fol. mod., fichero, …) — P73 note candidate',
}


def norm(s):
    return re.sub(r'\s+', ' ', s).strip().lower()


def load_export(path):
    delim = '\t' if path.endswith('.tsv') else ','
    with open(path, encoding='utf-8-sig') as f:
        return [(r['value'], int(r['freq'])) for r in csv.DictReader(f, delimiter=delim)]


def load_june(path):
    with open(path, encoding='utf-8') as f:
        return {r['value']: int(r['freq']) for r in csv.DictReader(f, delimiter='\t')}


def load_mapping(path):
    """
    Return (by_basis, by_key):
      by_basis: norm(basis) -> best row (vetted+valid QID preferred)
      by_key:   norm(key)   -> (qid, label) from vetted rows with valid QID
    """
    def rank(r):
        vetted = is_vetted(r['vetted'])
        return (vetted and bool(r['_qid']), vetted, bool(r['_qid']))

    by_basis, by_key = {}, {}
    with open(path, encoding='utf-8') as f:
        for r in csv.DictReader(f, delimiter='\t'):
            raw_qid = (r.get('qid') or '').strip()
            r['_qid'] = raw_qid if QID_RE.match(raw_qid) else extract_qid(raw_qid)
            r['_raw_qid'] = raw_qid
            b = norm(strip_html(r['basis']))
            if b not in by_basis or rank(r) > rank(by_basis[b]):
                by_basis[b] = r
            if is_vetted(r['vetted']) and r['_qid'] and not r.get('loc', '').strip():
                by_key.setdefault(norm(r['key']), (r['_qid'], r.get('label', '')))
    return by_basis, by_key


def load_done_counts(path):
    counts = Counter()
    if os.path.exists(path):
        with open(path, encoding='utf-8') as f:
            for line in f:
                try:
                    counts[norm(strip_html(json.loads(line)['basis']))] += 1
                except (ValueError, KeyError):
                    pass
    return counts


def classify(value, freq, june, by_basis, by_key, done):
    pp = preprocess(value)
    b = norm(pp['basis'])
    row = {
        'freq': freq,
        'freq_june': june.get(value, ''),
        'basis': value,
        'key': pp['key'],
        'loc': pp['loc'],
        'loc_type': pp['loc_type'],
        'qid': '', 'label': '', 'vetted': '', 'match_type': '', 'comment': '',
        'converted': done.get(b, 0) or '',
    }
    m = by_basis.get(b)
    if m:
        row.update(key=m['key'], loc=m.get('loc', ''), loc_type=m.get('loc_type', ''),
                   qid=m['_qid'], label=m.get('label', ''), vetted=m['vetted'],
                   match_type=m.get('match_type', ''), comment=m.get('comment', ''))

    if pp['preproc_type'] == 'excluded' or row['match_type'] == 'excluded':
        status = 'excluded'
    elif pp['preproc_type'] == 'compound':
        status = 'compound'
    elif m and is_vetted(m['vetted']):
        if not m['_qid'] or m['_qid'] != m['_raw_qid']:
            status = 'vetted_bad_qid'
            row['comment'] = (row['comment'] + f' [QID cell: {m["_raw_qid"]!r}]').strip()
        else:
            status = 'pending_locator' if row['loc'].strip() else 'pending_run'
    elif m:
        status = 'suggested' if m['_qid'] else 'unmatched'
    else:
        hit = by_key.get(norm(pp['key']))
        if hit:
            status = 'new_key_match'
            row['qid'], row['label'] = hit
        else:
            status = 'new'
    row['status'] = status
    return row


COLUMNS = ['freq', 'freq_june', 'basis', 'key', 'loc', 'loc_type', 'status',
           'qid', 'label', 'vetted', 'match_type', 'converted', 'comment']

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill('solid', fgColor='DDDDDD')
LINK_FONT = Font(color='0563C1', underline='single')


def write_sheet(ws, columns, rows, widths=None, qid_col=None):
    ws.append(columns)
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    for r in rows:
        ws.append([r.get(c, '') for c in columns])
        if qid_col and r.get(qid_col):
            cell = ws.cell(row=ws.max_row, column=columns.index(qid_col) + 1)
            cell.hyperlink = FG_ITEM + r[qid_col]
            cell.font = LINK_FONT
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('export', help='query-service export (value,freq), .csv or .tsv')
    ap.add_argument('--out', default=OUT_XLSX)
    args = ap.parse_args()

    current = load_export(args.export)
    june = load_june(JUNE_TSV)
    by_basis, by_key = load_mapping(MAPPING_TSV)
    done = load_done_counts(PROGRESS)

    rows = [classify(v, f, june, by_basis, by_key, done) for v, f in current]
    rows.sort(key=lambda r: -r['freq'])

    wb = Workbook()
    ws = wb.active
    ws.title = 'bases'
    write_sheet(ws, COLUMNS, rows, [7, 9, 40, 30, 12, 10, 16, 11, 35, 7, 12, 10, 40], 'qid')

    # Summary
    n, f = Counter(), Counter()
    for r in rows:
        n[r['status']] += 1
        f[r['status']] += r['freq']
    summary = [{'status': s, 'strings': n[s], 'usages': f[s], 'meaning': STATUS_NOTE[s]}
               for s in STATUS_ORDER]
    summary.append({'status': 'TOTAL', 'strings': sum(n.values()), 'usages': sum(f.values())})
    write_sheet(wb.create_sheet('summary', 0), ['status', 'strings', 'usages', 'meaning'],
                summary, [18, 10, 10, 70])

    # Changes since June
    cur = dict(current)
    changes = []
    for v, fj in june.items():
        if v not in cur:
            changes.append({'basis': v, 'freq_june': fj, 'freq_now': 0, 'change': 'gone'})
        elif cur[v] != fj:
            changes.append({'basis': v, 'freq_june': fj, 'freq_now': cur[v],
                            'change': 'fewer' if cur[v] < fj else 'more'})
    for v, fn in current:
        if v not in june:
            changes.append({'basis': v, 'freq_june': 0, 'freq_now': fn, 'change': 'new'})
    for c in changes:
        c['delta'] = c['freq_now'] - c['freq_june']
        c['converted'] = done.get(norm(strip_html(c['basis'])), 0) or ''
    changes.sort(key=lambda c: (c['change'], -abs(c['delta'])))
    write_sheet(wb.create_sheet('since_june'),
                ['change', 'basis', 'freq_june', 'freq_now', 'delta', 'converted'],
                changes, [8, 45, 10, 10, 8, 10])

    wb.save(args.out)

    print(f'{len(rows)} strings, {sum(f.values())} usages -> {args.out}\n')
    for s in summary:
        print(f"  {s['status']:16} {s['strings']:6} {s['usages']:7}")
    print('\nsince June:', dict(Counter(c['change'] for c in changes)))


if __name__ == '__main__':
    main()
