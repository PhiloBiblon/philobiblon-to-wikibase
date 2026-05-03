"""
lookup_basis_keys.py — Phase 2 of P721 bootstrap: search FactGrid for each
unique key in the parsed TSV and write the final seed TSV.

Deduplicates by key before searching — one API call per unique key regardless
of how many rows share it.  Checks the legacy xlsx first so already-known QIDs
skip the API entirely.

Usage (from pb2wb/):
    python prop_migration/lookup_basis_keys.py
    python prop_migration/lookup_basis_keys.py --dry-run
    python prop_migration/lookup_basis_keys.py --parsed prop_migration/P721-parsed.tsv
    python prop_migration/lookup_basis_keys.py --out prop_migration/P721-P129-seed.tsv
    python prop_migration/lookup_basis_keys.py --xlsx prop_migration/Reference\\ sources.xlsx
    python prop_migration/lookup_basis_keys.py --no-xlsx
"""

import os
import sys
import re
import csv
import json
import argparse

dir_path = os.path.dirname(os.path.realpath(__file__))
parent_dir_path = os.path.dirname(dir_path)
sys.path.append(parent_dir_path)

from tqdm import tqdm
from prop_migration.generate_basis_mapping import (
    OUT_COLUMNS,
    api_search_one,
    out_row,
    vetted_value,
)

PARSED_TSV          = 'prop_migration/P721-parsed.tsv'
OUT_TSV             = 'prop_migration/P721-P129-seed.tsv'
XLSX_DEFAULT        = 'prop_migration/Reference sources.xlsx'
CHECKPOINT_DEFAULT  = 'prop_migration/lookup_checkpoint.json'
CHECKPOINT_EVERY    = 500

_QID_RE = re.compile(r'Q\d+')


def load_legacy_map(xlsx_path):
    """
    Read the legacy xlsx and return basis string → QID dict.
    Extracts QIDs from hyperlinks on the 'key' column cells.
    Returns empty dict if file is missing or openpyxl is unavailable.
    """
    try:
        import openpyxl
    except ImportError:
        print('Warning: openpyxl not installed — skipping legacy xlsx seed.')
        return {}

    if not os.path.exists(xlsx_path):
        print(f'Warning: legacy xlsx not found at {xlsx_path!r} — skipping.')
        return {}

    wb = openpyxl.load_workbook(xlsx_path)
    ws = None
    for name in wb.sheetnames:
        if 'beta' in name.lower() or 'base' in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    header = None
    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(max_row=5, values_only=True), start=1):
        if row and any(isinstance(c, str) and c.strip().lower() in ('basis', 'key') for c in row):
            header = [str(c).strip().lower() if c else '' for c in row]
            header_row_idx = i
            break

    if header is None:
        print('Warning: could not find header row in legacy xlsx — skipping.')
        return {}

    try:
        basis_col = header.index('basis')
        key_col   = header.index('key')
    except ValueError:
        print('Warning: legacy xlsx missing "basis" or "key" column — skipping.')
        return {}

    legacy_map = {}
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        basis_cell = row[basis_col]
        key_cell   = row[key_col]
        basis_val  = basis_cell.value
        if not basis_val:
            continue
        basis_str = str(basis_val).strip()
        qid = ''
        if key_cell.hyperlink and key_cell.hyperlink.target:
            m = _QID_RE.search(key_cell.hyperlink.target)
            if m:
                qid = m.group(0)
        if qid:
            legacy_map[basis_str] = qid

    print(f'Loaded {len(legacy_map)} QID mappings from legacy xlsx: {xlsx_path}')
    return legacy_map


def load_parsed(path):
    rows = []
    with open(path, encoding='utf-8') as f:
        for row in csv.DictReader(f, delimiter='\t'):
            rows.append(dict(row))
    return rows


def main():
    parser = argparse.ArgumentParser(
        description='Phase 2: search FactGrid for unique keys and write seed TSV')
    parser.add_argument('--parsed',  default=PARSED_TSV,
                        help=f'Parsed TSV from parse_basis_values.py (default: {PARSED_TSV})')
    parser.add_argument('--out',     default=OUT_TSV,
                        help=f'Output seed TSV (default: {OUT_TSV})')
    parser.add_argument('--dry-run', action='store_true',
                        help='Show what would be searched without calling the API')
    xlsx_group = parser.add_mutually_exclusive_group()
    xlsx_group.add_argument('--xlsx', default=XLSX_DEFAULT,
                            help=f'Legacy xlsx for QID pre-seeding (default: {XLSX_DEFAULT})')
    xlsx_group.add_argument('--no-xlsx', dest='xlsx', action='store_const', const=None,
                            help='Skip legacy xlsx seeding')
    parser.add_argument('--checkpoint', default=CHECKPOINT_DEFAULT,
                        help=f'Checkpoint JSON for resuming interrupted runs (default: {CHECKPOINT_DEFAULT})')
    parser.add_argument('--checkpoint-every', type=int, default=CHECKPOINT_EVERY,
                        help=f'Save checkpoint every N keys (default: {CHECKPOINT_EVERY})')
    args = parser.parse_args()

    if not os.path.exists(args.parsed):
        print(f'Error: {args.parsed} not found.')
        print('Run first:  python prop_migration/parse_basis_values.py')
        sys.exit(1)

    legacy_map = load_legacy_map(args.xlsx) if args.xlsx else {}

    parsed = load_parsed(args.parsed)
    print(f'Loaded {len(parsed)} rows from {args.parsed}')

    excluded  = [r for r in parsed if r.get('preproc_type') == 'excluded']
    compound  = [r for r in parsed if r.get('preproc_type') == 'compound']
    to_search = [r for r in parsed if r.get('preproc_type') == '']

    # Deduplicate keys, preserving first-seen order (highest frequency first
    # because parse_basis_values.py writes in frequency order)
    seen = set()
    unique_keys = []
    for r in to_search:
        k = r['key']
        if k not in seen:
            seen.add(k)
            unique_keys.append(k)

    print(f'  excluded : {len(excluded)}  compound : {len(compound)}  to search: {len(to_search)}')
    print(f'  unique keys: {len(unique_keys)}  (dedup saves {len(to_search) - len(unique_keys)} API calls)')

    if args.dry_run:
        legacy_hits = sum(1 for k in unique_keys if k in legacy_map)
        print(f'  legacy hits: {legacy_hits}  API calls needed: {len(unique_keys) - legacy_hits}')
        print('(dry-run: no API calls made, no output written)')
        return

    # --- Load checkpoint if present ---
    checkpoint_path = args.checkpoint
    key_map = {}
    if os.path.exists(checkpoint_path):
        with open(checkpoint_path, encoding='utf-8') as f:
            raw = json.load(f)
        key_map = {k: tuple(v) for k, v in raw.items()}
        print(f'Resumed from checkpoint: {len(key_map)} keys already done '
              f'({checkpoint_path})')

    # --- Build key → (qid, label, match_type) map ---
    matched = legacy = 0
    keys_to_do = [k for k in unique_keys if k not in key_map]
    print(f'  keys to search: {len(keys_to_do)}  (skipping {len(unique_keys) - len(keys_to_do)} from checkpoint)')
    with tqdm(keys_to_do, unit='key') as bar:
        for i, key in enumerate(bar, start=1):
            if key in legacy_map:
                key_map[key] = (legacy_map[key], '', 'legacy')
                legacy += 1
            else:
                qid, label, _, subtype = api_search_one(key)
                if qid:
                    mtype = f'api_{subtype}' if subtype in ('label', 'alias') else 'api_fuzzy'
                    matched += 1
                else:
                    qid = label = ''
                    mtype = ''  # resolved per-row below (llm_pending vs none)
                key_map[key] = (qid, label, mtype)
            bar.set_postfix(matched=matched, legacy=legacy)
            if i % args.checkpoint_every == 0:
                with open(checkpoint_path, 'w', encoding='utf-8') as f:
                    json.dump({k: list(v) for k, v in key_map.items()}, f)

    # --- Assemble output rows ---
    out_rows = []

    for r in excluded:
        out_rows.append(out_row(r['freq'], r['basis'], r['key'],
                                r['loc'], r['loc_type'], 'excluded', '', ''))
    for r in compound:
        out_rows.append(out_row(r['freq'], r['basis'], r['key'],
                                r['loc'], r['loc_type'], 'compound', '', ''))
    for r in to_search:
        qid, label, mtype = key_map.get(r['key'], ('', '', ''))
        if not qid and not mtype:
            pp = r.get('parse_pattern', '')
            if pp == 'raw':
                mtype = 'llm_pending'
            elif pp == 'shelfmark':
                mtype = 'shelfmark'
            else:
                mtype = 'none'
        out_rows.append(out_row(r['freq'], r['basis'], r['key'],
                                r['loc'], r['loc_type'], mtype, qid, label,
                                vetted=vetted_value(mtype),
                                parse_pattern=r.get('parse_pattern', '')))

    out_rows.sort(key=lambda r: -(int(r['freq']) if str(r['freq']).isdigit() else 0))

    os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
    with open(args.out, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=OUT_COLUMNS, delimiter='\t')
        writer.writeheader()
        writer.writerows(out_rows)

    if os.path.exists(checkpoint_path):
        os.remove(checkpoint_path)
        print(f'Checkpoint removed: {checkpoint_path}')

    from collections import Counter
    mc = Counter(r['match_type'] for r in out_rows)
    print(f'\nWritten: {args.out}  ({len(out_rows)} rows)')
    print(f'  excluded    : {mc["excluded"]}')
    print(f'  compound    : {mc["compound"]}')
    print(f'  legacy      : {mc["legacy"]}')
    print(f'  api_label   : {mc["api_label"]}')
    print(f'  api_alias   : {mc["api_alias"]}')
    print(f'  api_fuzzy   : {mc["api_fuzzy"]}')
    print(f'  none        : {mc["none"]}')
    print(f'  llm_pending : {mc["llm_pending"]}')
    print(f'\nNext step: upload {args.out} to Google Sheets as the initial P721-P129 sheet.')


if __name__ == '__main__':
    main()
