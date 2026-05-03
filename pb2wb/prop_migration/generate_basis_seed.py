"""
generate_basis_seed.py — bootstrap script for P721 basis/reference migration.

Run ONCE to generate the initial seed TSV that becomes the Google Sheet.
After uploading the seed to Google Sheets, use generate_basis_mapping.py
(read: fill-gaps) for all subsequent iterations.

Reads pre-fetched P721 values from fetch_basis_values.py, pre-processes each,
optionally seeds QIDs from a legacy xlsx file, searches FactGrid for a P129
item candidate for remaining rows, and writes a TSV ready to upload as the
initial Google Sheet.

Usage (from pb2wb/):
    # Step 0 (once): fetch raw values from FactGrid
    python prop_migration/fetch_basis_values.py
    python prop_migration/fetch_basis_values.py --limit 500

    # Step 1: generate the seed TSV from the fetched values
    python prop_migration/generate_basis_seed.py
    python prop_migration/generate_basis_seed.py --dry-run
    python prop_migration/generate_basis_seed.py --limit 200
    python prop_migration/generate_basis_seed.py --values prop_migration/P721-raw-values.tsv
    python prop_migration/generate_basis_seed.py --out prop_migration/P721-P129-seed.tsv
    python prop_migration/generate_basis_seed.py --xlsx prop_migration/Reference\ sources.xlsx
    python prop_migration/generate_basis_seed.py --no-xlsx
"""

import os
import re
import sys
import csv
import argparse

dir_path = os.path.dirname(os.path.realpath(__file__))
parent_dir_path = os.path.dirname(dir_path)
sys.path.append(parent_dir_path)

from tqdm import tqdm
from prop_migration.generate_basis_mapping import (
    OUT_COLUMNS,
    api_search_one,
    out_row,
    preprocess,
    vetted_value,
)

OUT_TSV      = 'prop_migration/P721-P129-seed.tsv'
VALUES_TSV   = 'prop_migration/P721-raw-values.tsv'
XLSX_DEFAULT = 'prop_migration/Reference sources.xlsx'

_QID_RE = re.compile(r'Q\d+')


def load_legacy_map(xlsx_path):
    """
    Read the legacy xlsx file and return a dict mapping basis string → QID.
    Extracts QIDs from cell hyperlinks (cell.hyperlink.target) in the 'key'
    column — the cell value may be plain text, so values_only=True would miss them.
    Returns empty dict if the file is not found or openpyxl is unavailable.
    """
    try:
        import openpyxl
    except ImportError:
        print('Warning: openpyxl not installed — skipping legacy xlsx seed.')
        return {}

    if not os.path.exists(xlsx_path):
        print(f'Warning: legacy xlsx not found at {xlsx_path!r} — skipping.')
        return {}

    wb = openpyxl.load_workbook(xlsx_path)
    # Use first sheet that looks right, or fallback to first sheet
    sheet_names = wb.sheetnames
    ws = None
    for name in sheet_names:
        if 'beta' in name.lower() or 'base' in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb[sheet_names[0]]

    # Find header row to locate 'basis' and 'key' columns
    header = None
    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(max_row=5, values_only=True), start=1):
        if row and any(isinstance(c, str) and c.strip().lower() in ('basis', 'key') for c in row):
            header = [str(c).strip().lower() if c else '' for c in row]
            header_row_idx = i
            break

    if header is None:
        print('Warning: could not find header row in legacy xlsx — skipping.')
        return {}

    try:
        basis_col = header.index('basis')
        key_col   = header.index('key')
    except ValueError:
        print('Warning: legacy xlsx missing "basis" or "key" column — skipping.')
        return {}

    legacy_map = {}
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        basis_cell = row[basis_col]
        key_cell   = row[key_col]
        basis_val  = basis_cell.value
        if not basis_val:
            continue
        basis_str = str(basis_val).strip()
        # Extract QID from hyperlink on the key cell
        qid = ''
        if key_cell.hyperlink and key_cell.hyperlink.target:
            m = _QID_RE.search(key_cell.hyperlink.target)
            if m:
                qid = m.group(0)
        if qid:
            legacy_map[basis_str] = qid

    print(f'Loaded {len(legacy_map)} QID mappings from legacy xlsx: {xlsx_path}')
    return legacy_map


def load_values(tsv_path, limit=None):
    """Read (value, freq) pairs from the TSV written by fetch_basis_values.py."""
    rows = []
    with open(tsv_path, encoding='utf-8') as f:
        for row in csv.DictReader(f, delimiter='\t'):
            rows.append((row['value'], int(row['freq'])))
            if limit and len(rows) >= limit:
                break
    return rows


def main():
    parser = argparse.ArgumentParser(
        description='Bootstrap: pre-process P721 values and write seed TSV for Google Sheets')
    parser.add_argument('--values', default=VALUES_TSV,
                        help=f'Raw values TSV from fetch_basis_values.py (default: {VALUES_TSV})')
    parser.add_argument('--out',    default=OUT_TSV,
                        help=f'Output seed TSV (default: {OUT_TSV})')
    parser.add_argument('--limit',  type=int, default=None,
                        help='Process only the top N values from the input (default: all)')
    parser.add_argument('--dry-run', action='store_true',
                        help='Show what would be searched without calling the API')
    xlsx_group = parser.add_mutually_exclusive_group()
    xlsx_group.add_argument('--xlsx', default=None,
                            help='Legacy xlsx for pre-seeding QIDs (disabled by default; '
                                 'use reference_source.gold_seed.tsv via generate_basis_mapping.py instead)')
    xlsx_group.add_argument('--no-xlsx', dest='xlsx', action='store_const', const=None,
                            help='Skip legacy xlsx seeding (default)')
    args = parser.parse_args()

    legacy_map = load_legacy_map(args.xlsx) if args.xlsx else {}

    if not os.path.exists(args.values):
        print(f'Error: values file not found: {args.values}')
        print('Run first:  python prop_migration/fetch_basis_values.py')
        sys.exit(1)
    print(f'Reading values from: {args.values}')
    values = load_values(args.values, args.limit)
    limit_msg = f'top {args.limit}' if args.limit else 'all'
    print(f'Loaded {len(values)} values ({limit_msg}).')

    preprocessed = [{'freq': freq, 'raw': raw, 'pp': preprocess(raw)}
                    for raw, freq in values]

    excluded  = [r for r in preprocessed if r['pp']['preproc_type'] == 'excluded']
    compound  = [r for r in preprocessed if r['pp']['preproc_type'] == 'compound']
    to_search = [r for r in preprocessed if r['pp']['preproc_type'] == '']

    print(f'  excluded : {len(excluded)}  compound : {len(compound)}  to search: {len(to_search)}')

    if args.dry_run:
        for r in to_search:
            pp = r['pp']
            print(f"  [{pp['parse_pattern']:14s}] key={pp['key']!r:30s}  "
                  f"loc={pp['loc']!r}  ({r['freq']}×)")
        print(f'\n(dry-run: {len(excluded)} excluded, {len(compound)} compound, '
              f'{len(to_search)} would be searched — no API calls made, no output written)')
        return

    out_rows = []

    for r in excluded:
        pp = r['pp']
        out_rows.append(out_row(r['freq'], pp['basis'], pp['key'],
                                pp['loc'], pp['loc_type'], 'excluded', '', ''))
    for r in compound:
        pp = r['pp']
        out_rows.append(out_row(r['freq'], pp['basis'], pp['key'],
                                pp['loc'], pp['loc_type'], 'compound', '', ''))

    matched = legacy = 0
    with tqdm(to_search, unit='basis') as bar:
        for r in bar:
            pp = r['pp']
            basis_str = pp['basis']
            if basis_str in legacy_map:
                qid   = legacy_map[basis_str]
                mtype = 'legacy'
                label = ''
                legacy += 1
            else:
                qid, label, _, subtype = api_search_one(pp['key'])
                if qid:
                    mtype = f'api_{subtype}' if subtype in ('label', 'alias') else 'api_fuzzy'
                    matched += 1
                else:
                    mtype = 'llm_pending' if pp['parse_pattern'] == 'raw' else 'none'
                    qid = label = ''
            out_rows.append(out_row(r['freq'], pp['basis'], pp['key'],
                                    pp['loc'], pp['loc_type'], mtype, qid, label,
                                    vetted=vetted_value(mtype)))
            bar.set_postfix(matched=matched, legacy=legacy)

    out_rows.sort(key=lambda r: -int(r['freq']))

    os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
    with open(args.out, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=OUT_COLUMNS, delimiter='\t')
        writer.writeheader()
        writer.writerows(out_rows)

    from collections import Counter
    mc = Counter(r['match_type'] for r in out_rows)
    print(f'\nWritten: {args.out}  ({len(out_rows)} rows)')
    print(f'  excluded    : {mc["excluded"]}')
    print(f'  compound    : {mc["compound"]}')
    print(f'  legacy      : {mc["legacy"]}')
    print(f'  api_label   : {mc["api_label"]}')
    print(f'  api_alias   : {mc["api_alias"]}')
    print(f'  api_fuzzy   : {mc["api_fuzzy"]}')
    print(f'  none        : {mc["none"]}')
    print(f'  llm_pending : {mc["llm_pending"]}')
    print(f'\nNext step: upload {args.out} to Google Sheets as the initial P721-P129 sheet.')


if __name__ == '__main__':
    main()
